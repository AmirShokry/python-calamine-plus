"""openpyxl's utilities implemented in Rust, checked against openpyxl itself."""

from __future__ import annotations

import datetime
import random

import pytest
from python_calamine_plus.styles.numbers import (
    BUILTIN_FORMATS,
    is_date_format,
    is_timedelta_format,
)
from python_calamine_plus.utils import (
    CellCoordinatesException,
    absolute_coordinate,
    cols_from_range,
    column_index_from_string,
    coordinate_from_row_col,
    coordinate_from_string,
    coordinate_from_tuple,
    coordinate_to_tuple,
    get_column_interval,
    get_column_letter,
    quote_sheetname,
    range_boundaries,
    range_to_tuple,
    rows_from_range,
    split_print_areas,
)
from python_calamine_plus.utils.datetime import (
    CALENDAR_MAC_1904,
    MAC_EPOCH,
    WINDOWS_EPOCH,
    from_excel,
    from_ISO8601,
    serial_to_datetime,
    to_excel,
)

ox = pytest.importorskip("openpyxl")
from openpyxl.styles import numbers as ox_numbers  # noqa: E402
from openpyxl.utils import cell as ox_cell  # noqa: E402
from openpyxl.utils import datetime as ox_dt  # noqa: E402


def _same(ours, theirs, *args):
    """Same result, or the same exception type."""
    try:
        expected = theirs(*args)
    except Exception as e:  # noqa: BLE001
        with pytest.raises(Exception) as info:  # noqa: B017
            ours(*args)
        # Same kind of error (our exception classes are not openpyxl's).
        assert type(info.value).__name__ == type(e).__name__ or (
            isinstance(e, ValueError) and isinstance(info.value, ValueError)
        ), (args, e, info.value)
        return
    got = ours(*args)
    if hasattr(expected, "__next__"):
        expected, got = list(expected), list(got)
    assert got == expected, args


def test_columns():
    for i in range(1, 18279):
        assert get_column_letter(i) == ox_cell.get_column_letter(i)
        assert column_index_from_string(ox_cell.get_column_letter(i)) == i
    for bad in (0, -1, 18279):
        _same(get_column_letter, ox_cell.get_column_letter, bad)
    for s in ("a", "zz", "Zz", "", "AAAA", "A1", "$A", "ZZZ"):
        _same(column_index_from_string, ox_cell.column_index_from_string, s)


COORDS = [
    "A1",
    "$B$12",
    "b7",
    "ZZZ1048576",
    "AAAA1",
    "A0",
    "1A",
    "A",
    "",
    "A1:B2",
    "$A1",
]
RANGES = [
    "A1",
    "A1:C3",
    "$A$1:$C$3",
    "C3:A1",
    "A:C",
    "1:3",
    "$A:$C",
    "$1:$3",
    "A1:C",
    "A:3",
    "",
    "A1:",
    "AAAA1",
    "A1:B2:C3",
    "a1:c3",
    "A$1",
]


def test_coordinates():
    for c in COORDS:
        _same(coordinate_from_string, ox_cell.coordinate_from_string, c)
        _same(absolute_coordinate, ox_cell.absolute_coordinate, c)
    for c in ["A1", "B12", "zz9", "AB100", "A 1"]:
        _same(coordinate_to_tuple, ox_cell.coordinate_to_tuple, c)
    with pytest.raises(CellCoordinatesException):
        coordinate_from_string("A0")
    assert coordinate_from_row_col(7, 3) == "C7"
    assert coordinate_from_tuple((7, 3)) == "C7"
    assert coordinate_from_tuple(coordinate_to_tuple("XFD1048576")) == "XFD1048576"
    with pytest.raises(ValueError):
        coordinate_from_row_col(0, 1)
    with pytest.raises(ValueError):
        coordinate_from_row_col(1, 18279)


def test_ranges():
    for r in RANGES:
        _same(range_boundaries, ox_cell.range_boundaries, r)
        _same(absolute_coordinate, ox_cell.absolute_coordinate, r)
    for r in ["A1:C3", "B2", "C3:A1"]:
        _same(rows_from_range, ox_cell.rows_from_range, r)
        _same(cols_from_range, ox_cell.cols_from_range, r)
    for r in [
        "Sheet1!A1:B2",
        "'My Sheet'!$A$1:$B$2",
        "'It''s'!A1",
        "Sheet1!A:B",
        "Sheet1!1:3",
        "A1:B2",
        "Sheet 1!A1",
    ]:
        _same(range_to_tuple, ox_cell.range_to_tuple, r)
    for name in ("Data", "My Sheet", "It's"):
        assert quote_sheetname(name) == ox_cell.quote_sheetname(name)
    assert get_column_interval("A", "D") == ox_cell.get_column_interval("A", "D")
    assert get_column_interval(2, 5) == ox_cell.get_column_interval(2, 5)


def test_split_print_areas():
    assert split_print_areas("'Data'!$A$1:$C$10,'Data'!$E$1:$F$2") == [
        "A1:C10",
        "E1:F2",
    ]
    assert split_print_areas("Sheet1!$B$2") == ["B2"]
    assert split_print_areas("'a,b'!$A$1:$B$2") == ["A1:B2"]
    assert split_print_areas("Sheet1!$1:$3,Sheet1!$A:$B") == ["1:3", "A:B"]
    for dynamic in ("OFFSET(Sheet1!$A$1,0,0,10,2)", "Sheet1!MyRange", "", "Sheet1!"):
        with pytest.raises(ValueError):
            split_print_areas(dynamic)


def test_number_formats():
    formats = list(ox_numbers.BUILTIN_FORMATS.values()) + [
        "yyyy-mm-dd",
        "[h]:mm:ss",
        "[HH]:MM",
        "[mm]:ss",
        "[ss].00",
        '"days" 0',
        "[Red]0.00",
        "[$-409]mmmm d, yyyy",
        "_d0",
        "\\d0",
        '0.00;[Red]-0.00;"zero"',
        "General",
        "@",
        "[h]",
        "d/m/y;@",
        '"d"',
        "[$€-2] #,##0.00",
        "",
    ]
    for f in formats:
        assert is_date_format(f) == ox_numbers.is_date_format(f), f
        assert is_timedelta_format(f) == ox_numbers.is_timedelta_format(f), f
    assert is_date_format(None) is False
    assert BUILTIN_FORMATS == ox_numbers.BUILTIN_FORMATS


def test_from_excel_matches_openpyxl():
    rng = random.Random(1234)
    values = [
        0,
        0.5,
        0.999999999,
        1,
        59,
        60,
        61,
        60.5,
        45351,
        45351.123456,
        -1,
        -0.5,
        2958465.99,
    ]
    values += [rng.uniform(-100, 3_000_000) for _ in range(3000)]
    values += [rng.randint(0, 3_000_000) for _ in range(500)]
    values += [round(rng.uniform(0, 1), rng.randint(1, 12)) for _ in range(1000)]
    for epoch in (WINDOWS_EPOCH, MAC_EPOCH):
        for td in (False, True):
            for v in values:
                try:
                    expected = ox_dt.from_excel(v, epoch, timedelta=td)
                except (OverflowError, ValueError):
                    with pytest.raises((OverflowError, ValueError)):
                        from_excel(v, epoch, td)
                    continue
                assert from_excel(v, epoch, td) == expected, (v, epoch, td)
    assert from_excel(None) is None
    assert CALENDAR_MAC_1904 == MAC_EPOCH == ox_dt.MAC_EPOCH


def test_to_excel_matches_openpyxl():
    cases = [
        datetime.datetime(2024, 2, 29, 13, 45, 1, 500000),
        datetime.datetime(1900, 1, 1),
        datetime.datetime(1900, 2, 28),
        datetime.datetime(1900, 3, 1),
        datetime.datetime(1899, 12, 30),
        datetime.datetime(1899, 12, 29, 12),
        datetime.date(2020, 5, 17),
        datetime.time(18, 30, 15, 250000),
        datetime.timedelta(days=2, hours=3),
        datetime.timedelta(seconds=-90),
    ]
    for epoch in (WINDOWS_EPOCH, MAC_EPOCH):
        for c in cases:
            assert to_excel(c, epoch) == ox_dt.to_excel(c, epoch), c


def test_iso8601_matches_openpyxl():
    for s in [
        "2024-01-31",
        "2024-01-31T10:30:00",
        "2024-01-31T10:30:00Z",
        "2024-01-31T10:30:00.123Z",
        "2024-01-31T10:30",
        "10:30:15",
        "T10:30:15.5",
        "PT10H30M",
        "PT1.5S",
        "PT255H10M10S",
        "",
        "2024-01-31 10:30:00",
    ]:
        _same(from_ISO8601, ox_dt.from_ISO8601, s)
    with pytest.raises(ValueError):
        from_ISO8601("not a date")


def test_serial_to_datetime():
    assert serial_to_datetime(45351) == datetime.datetime(2024, 2, 29)
    assert serial_to_datetime(45351.5) == datetime.datetime(2024, 2, 29, 12)
    assert serial_to_datetime(0.25) == datetime.datetime(1899, 12, 30, 6)
    assert serial_to_datetime(59) == datetime.datetime(1900, 2, 28)
    assert serial_to_datetime(61) == datetime.datetime(1900, 3, 1)
    assert serial_to_datetime(0, epoch_1904=True) == datetime.datetime(1904, 1, 1)
    with pytest.raises(ValueError):
        serial_to_datetime(1e12)
