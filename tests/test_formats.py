"""xlsx, xlsm (macro-enabled) and the xltx / xltm templates all stream identically."""

import zipfile
from pathlib import Path

import pytest
from python_calamine_plus import CalamineWorkbook, StreamingNotSupported

openpyxl = pytest.importorskip("openpyxl")

PATH = Path(__file__).parent / "data"

CONTENT_TYPES = {
    "xlsm": b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    "xltx": b"application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml",
    "xltm": b"application/vnd.ms-excel.template.macroEnabled.main+xml",
}
XLSX_TYPE = (
    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)


@pytest.fixture(scope="module")
def source(tmp_path_factory):
    from openpyxl.comments import Comment
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["name", "n", "f"])
    for i in range(1, 50):
        ws.append([f"r{i}", i, f"=B{i + 1}*2"])
    ws["A1"].font = Font(bold=True)
    ws["A2"].comment = Comment("first", "me")
    ws["C1"].comment = Comment("header", "me")
    ws.merge_cells("D1:E1")
    ws.freeze_panes = "A2"
    path = tmp_path_factory.mktemp("formats") / "book.xlsx"
    wb.save(path)
    return path


def convert(src, dst, kind):
    """Same package, different workbook content type (+ a VBA part for macro files)."""
    with (
        zipfile.ZipFile(src) as zin,
        zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                assert XLSX_TYPE in data
                data = data.replace(XLSX_TYPE, CONTENT_TYPES[kind])
            zout.writestr(item, data)
        if kind in ("xlsm", "xltm"):
            zout.writestr(
                "xl/vbaProject.bin", b"\xd0\xcf\x11\xe0 not a real VBA project"
            )


def snapshot(path):
    wb = CalamineWorkbook.from_path(path)
    out = [wb.sheet_names, wb.defined_names, wb.properties, wb.named_styles]
    for name in wb.sheet_names:
        out.append(list(wb.stream_sheet_by_name(name)))
        s = wb.stream_sheet_by_name(name, rich=True)
        rows = [
            [(c.value, c.formula, c.style, c.comment, c.merged_range) for c in r]
            for r in s
        ]
        out += [rows, s.sheet_settings, s.comments, s.merged_cell_ranges]
        out.append(wb.get_sheet_by_name(name).to_python())
    return out


@pytest.mark.parametrize("kind", ["xlsm", "xltx", "xltm"])
def test_same_output_as_xlsx(source, tmp_path, kind):
    dst = tmp_path / f"book.{kind}"
    convert(source, dst, kind)
    assert snapshot(dst) == snapshot(source)
    wb = CalamineWorkbook.from_path(dst)
    assert wb.has_macros is (kind in ("xlsm", "xltm"))
    assert wb.is_template is kind.startswith("xlt")
    # From bytes too (no extension to go by).
    with open(dst, "rb") as f:
        assert list(CalamineWorkbook.from_filelike(f).stream_sheet_by_index(0)) == list(
            wb.stream_sheet_by_index(0)
        )


def test_comments_in_sheet_order(source):
    s = CalamineWorkbook.from_path(source).stream_sheet_by_index(0, rich=True)
    assert list(s.comments) == [(0, 2), (1, 0)]


@pytest.mark.parametrize("filename", ["base.xlsb", "base.xls", "base.ods"])
def test_other_formats_not_streamable(filename):
    wb = CalamineWorkbook.from_path(PATH / filename)
    with pytest.raises(StreamingNotSupported, match="xlsx / xlsm"):
        wb.stream_sheet_by_index(0)
    assert wb.get_sheet_by_index(0).to_python()  # the full load still works
