"""`load_workbook(..., compatibility="openpyxl")`: parity with openpyxl (used as the
oracle), lifecycle, streaming and resource behavior."""

from __future__ import annotations

import datetime
import itertools
import warnings
from pathlib import Path

import pytest
from python_calamine_plus import (
    CalamineWorkbook,
    CompatibilityNotSupported,
    WorkbookClosed,
    load_workbook,
)

from .xlsx_builder import build, sheet

openpyxl = pytest.importorskip("openpyxl")
warnings.filterwarnings("ignore", module="openpyxl")

DATA = Path(__file__).parent / "data"
MODES = [(False, False), (False, True), (True, False), (True, True)]


# ---- oracle ------------------------------------------------------------------------

PLAIN = (
    type(None),
    bool,
    int,
    float,
    str,
    datetime.datetime,
    datetime.date,
    datetime.time,
    datetime.timedelta,
)


def _unset(v):
    # openpyxl returns the descriptor object for color attributes that are not set;
    # this profile returns None.
    return None if type(v).__module__.startswith("openpyxl.") else v


def _color(c):
    if c is None:
        return None
    return tuple(
        _unset(getattr(c, a))
        for a in ("rgb", "indexed", "theme", "auto", "tint", "type")
    )


def _side(s):
    return None if s is None else (s.style, _color(s.color))


def _style(c):
    f, a, b, p, fill = c.font, c.alignment, c.border, c.protection, c.fill
    if fill.tagname == "gradientFill":
        fill_key = (
            "gradient",
            fill.type,
            fill.degree,
            [(s.position, _color(s.color)) for s in fill.stop],
        )
    else:
        fill_key = (fill.patternType, _color(fill.fgColor), _color(fill.bgColor))
    return (
        c.number_format,
        (f.name, f.sz, f.b, f.i, f.u, f.strike, f.vertAlign, _color(f.color), f.family)
        + (f.charset, f.scheme, f.outline, f.shadow, f.condense, f.extend),
        fill_key,
        tuple(
            _side(getattr(b, n))
            for n in (
                "left",
                "right",
                "top",
                "bottom",
                "diagonal",
                "vertical",
                "horizontal",
            )
        )
        + (b.diagonalUp, b.diagonalDown, b.outline),
        (a.horizontal, a.vertical, a.textRotation, a.wrapText, a.shrinkToFit, a.indent)
        + (a.relativeIndent, a.justifyLastLine, a.readingOrder),
        (p.locked, p.hidden),
    )


def _value(v):
    name = type(v).__name__
    if name == "ArrayFormula":
        return ("ArrayFormula", v.ref, v.text)
    if name == "DataTableFormula":
        return ("DataTable", v.ref, v.ca, v.dt2D, v.dtr, v.r1, v.r2, v.del1, v.del2)
    assert isinstance(v, PLAIN), v
    return (type(v).__name__, v)


def _cell(c):
    kind = type(c).__name__
    out = [kind, _value(c.value), c.data_type]
    if kind == "EmptyCell":
        return out
    out += [c.coordinate, c.row, c.column, _style(c)]
    if kind in ("Cell", "ReadOnlyCell"):
        out.append(c.is_date)
    if kind == "Cell":
        out.append((c.comment.text, c.comment.author) if c.comment else None)
        h = c.hyperlink
        out.append((h.ref, h.target, h.location, h.tooltip, h.display) if h else None)
    return out


def _sheet_info(ws, read_only):
    out = [
        ws.title,
        ws.sheet_state,
        ws.min_row,
        ws.min_column,
        ws.max_row,
        ws.max_column,
    ]
    if read_only:
        try:
            out.append(ws.calculate_dimension())
        except ValueError as e:
            out.append(str(e))
        return out
    out += [
        ws.dimensions,
        str(ws.merged_cells),
        ws.freeze_panes,
        ws.print_area,
        ws.print_titles,
        ws.print_title_rows,
        ws.print_title_cols,
        sorted((k, v.value) for k, v in ws.defined_names.items()),
        sorted(
            (k, d.height, d.hidden, d.outlineLevel, d.customHeight)
            for k, d in ws.row_dimensions.items()
        ),
        sorted(
            (
                k,
                d.width,
                d.hidden,
                d.outlineLevel,
                d.min,
                d.max,
                d.bestFit,
                d.customWidth,
            )
            for k, d in ws.column_dimensions.items()
        ),
        ws.auto_filter.ref,
        _color(ws.sheet_properties.tabColor),
    ]
    return out


def assert_parity(path, read_only, data_only, **iter_kw):
    """Workbook, sheet, row and cell attributes equal openpyxl's."""
    ox = openpyxl.load_workbook(path, read_only=read_only, data_only=data_only)
    cp = load_workbook(
        path, compatibility="openpyxl", read_only=read_only, data_only=data_only
    )
    assert cp.sheetnames == ox.sheetnames
    assert [w.title for w in cp.worksheets] == [w.title for w in ox.worksheets]
    assert cp.active.title == ox.active.title
    assert cp.epoch == ox.epoch
    if not read_only:
        assert sorted((k, v.value) for k, v in cp.defined_names.items()) == sorted(
            (k, v.value) for k, v in ox.defined_names.items()
        )
    for ows in ox.worksheets:
        cws = cp[ows.title]
        assert _sheet_info(cws, read_only) == _sheet_info(ows, read_only), ows.title
        for values_only in (True, False):
            orows = list(ows.iter_rows(values_only=values_only, **iter_kw))
            crows = list(cws.iter_rows(values_only=values_only, **iter_kw))
            assert len(crows) == len(orows), (ows.title, values_only)
            for i, (orow, crow) in enumerate(zip(orows, crows)):
                assert type(crow) is tuple
                assert len(crow) == len(orow), (ows.title, i)
                for oc, cc in zip(orow, crow):
                    if values_only:
                        assert _value(cc) == _value(oc), (ows.title, i)
                    else:
                        assert _cell(cc) == _cell(oc), (ows.title, i)
    ox.close()
    cp.close()


# ---- fixtures ----------------------------------------------------------------------


@pytest.fixture(scope="module")
def styled_xlsx(tmp_path_factory):
    """Styles, merged cells (with borders), hyperlinks, comments, print settings,
    dimensions, defined names, hidden sheets: written by openpyxl."""
    from openpyxl.comments import Comment
    from openpyxl.styles import (
        Alignment,
        Border,
        Color,
        Font,
        GradientFill,
        PatternFill,
        Protection,
        Side,
    )
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Header"
    ws["A1"].font = Font(
        name="Arial", size=14, bold=True, italic=True, color="FFFF0000"
    )
    ws["A1"].fill = PatternFill("solid", fgColor="FFFFFF00")
    ws["A1"].border = Border(
        top=Side("thick"), left=Side("thin", color="FF0000FF"), bottom=Side("double")
    )
    ws["A1"].alignment = Alignment(horizontal="center", wrap_text=True, indent=2)
    ws.merge_cells("A1:C2")
    ws["E1"] = "end-styled"
    ws.merge_cells("E1:F3")
    ws["F3"].border = Border(right=Side("medium"), bottom=Side("dashed"))
    ws["A4"] = 1
    ws["B4"] = 2.5
    ws["C4"] = datetime.datetime(2024, 2, 29, 13, 45)
    ws["D4"] = datetime.date(2020, 1, 1)
    ws["E4"] = datetime.time(6, 30)
    ws["F4"] = datetime.timedelta(hours=30)
    ws["G4"] = True
    ws["H4"] = "#N/A"
    ws["A5"] = "=SUM(A4:B4)"
    ws["B5"] = '=CONCATENATE("a", "b")'
    ws["C5"] = "theme"
    ws["C5"].font = Font(
        color=Color(theme=4, tint=-0.25), underline="double", strike=True
    )
    ws["C5"].fill = GradientFill(degree=45, stop=("FF000000", "FFFFFFFF"))
    ws["D5"] = "indexed"
    ws["D5"].fill = PatternFill("darkGrid", fgColor=Color(indexed=10), bgColor="00FF00")
    ws["D5"].protection = Protection(locked=False, hidden=True)
    ws["D5"].font = Font(vertAlign="superscript", outline=False)
    ws["A7"] = "linked"
    ws["A7"].hyperlink = "https://example.com/a?x=1&y=2"  # escaped in the .rels part
    ws["B7"].hyperlink = "https://example.com/empty"  # fills the empty cell
    ws["C7"].hyperlink = "#Data!A1"
    ws["A8"] = "noted"
    ws["A8"].comment = Comment("first", "Ann")
    ws["J12"].comment = Comment("far away", "Bo")  # extends the bounds
    ws["D10"].fill = PatternFill("solid", fgColor="FF00FF00")  # styled, no value
    ws["A9"] = ""  # openpyxl writes no value for an empty string
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions.group("D", "F", hidden=True, outline_level=1)
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[20].hidden = True  # below the cells: not part of the bounds
    ws.freeze_panes = "B2"
    ws.print_area = ["A1:C10", "E1:F2"]
    ws.print_title_rows = "1:2"
    ws.print_title_cols = "A:B"
    ws.auto_filter.ref = "A4:H5"
    ws.sheet_properties.tabColor = "FF123456"
    ws.defined_names["local_name"] = DefinedName("local_name", attr_text="Data!$A$4")
    wb.defined_names["global_name"] = DefinedName(
        "global_name", attr_text="Data!$B$4:$C$4"
    )

    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["B2"] = "x"
    wb.create_sheet("Empty")
    wb.active = 0
    path = tmp_path_factory.mktemp("compat") / "styled.xlsx"
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def raw_xlsx(tmp_path_factory):
    """Value and formula cases openpyxl cannot write, as Excel stores them."""
    rows = (
        # literal ints/floats, dates (style 1), timedelta (2), time (3), "days" text format (5)
        '<row r="1"><c r="A1"><v>5</v></c><c r="B1"><v>5.0</v></c><c r="C1"><v>1e3</v></c>'
        '<c r="D1" s="1"><v>45351</v></c><c r="E1" s="2"><v>1.25</v></c>'
        '<c r="F1" s="3"><v>0.5</v></c><c r="G1" s="5"><v>12</v></c>'
        '<c r="H1" s="7"><v>45351.5</v></c><c r="I1" s="1"><v>59</v></c>'
        '<c r="J1" s="1"><v>61</v></c><c r="K1" s="8"><v>-3</v></c></row>'
        # strings: shared, empty shared, inline, empty inline, formula string result
        '<row r="2"><c r="A2" t="s"><v>0</v></c><c r="B2" t="s"><v>1</v></c>'
        '<c r="C2" t="inlineStr"><is><t>inline</t></is></c>'
        '<c r="D2" t="inlineStr"><is><t></t></is></c>'
        '<c r="E2" t="str"><f>"x"&amp;"y"</f><v>xy</v></c>'
        '<c r="F2" t="b"><v>0</v></c><c r="G2" t="e"><v>#DIV/0!</v></c>'
        '<c r="H2" t="d"><v>2024-01-31T10:30:00</v></c><c r="I2" t="d"><v>2024-01-31</v></c>'
        '<c r="J2" s="6"/></row>'
        # shared formula with cached values, array formula, formula without cached value
        '<row r="4"><c r="A4"><f t="shared" ref="A4:C4" si="0">A1+1</f><v>6</v></c>'
        '<c r="B4"><f t="shared" si="0"/><v>6</v></c><c r="C4"><f t="shared" si="0"/><v>1001</v></c>'
        '<c r="D4"><f t="array" ref="D4:D5">A1:A2*2</f><v>10</v></c>'
        '<c r="E4"><f>A1/0</f></c><c r="F4" t="e"><f>1/0</f><v>#DIV/0!</v></c>'
        '<c r="G4" t="str"><f>""</f><v></v></c></row>'
        '<row r="5"><c r="D5"><v>4</v></c></row>'
        # data table formula
        '<row r="6"><c r="B6"><f t="dataTable" ref="B6:B7" dt2D="0" dtr="1" r1="A1"/><v>3</v></c></row>'
        # cells without r attributes
        '<row r="8"><c><v>1</v></c><c><v>2</v></c><c r="E8"><v>5</v></c><c><v>6</v></c></row>'
        # a row without cells at the end (hidden)
        '<row r="12" hidden="1"/>'
    )
    path = tmp_path_factory.mktemp("compat") / "raw.xlsx"
    return build(
        path,
        {
            "Values": sheet(rows, dimension="A1:K12"),
            "NoDim": sheet(
                '<row r="2"><c r="B2"><v>1</v></c><c r="D2"><v>2</v></c></row>'
                '<row r="4"><c r="C4"><v>3</v></c></row>',
                dimension=None,
            ),
            # A link on a merged placeholder is bound to the range's top-left cell.
            "Merged": sheet(
                '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1"><v>9</v></c></row>'
                '<row r="2"><c r="A2"/><c r="B2"/></row>',
                dimension="A1:B2",
                extra='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>'
                '<hyperlinks><hyperlink ref="B2" location="Values!A1" display="go"/>'
                '<hyperlink ref="C3" location="Values!B1"/></hyperlinks>',
            ),
            "Sparse": sheet(
                '<row r="1"><c r="A1"><v>1</v></c></row>'
                '<row r="1048576"><c r="XFD1048576"><v>2</v></c></row>',
                dimension="A1:XFD1048576",
            ),
        },
        shared_strings=["shared", ""],
    )


@pytest.fixture(scope="module")
def epoch_1904_xlsx(tmp_path_factory):
    rows = (
        '<row r="1"><c r="A1" s="1"><v>0</v></c><c r="B1" s="1"><v>1</v></c>'
        '<c r="C1" s="7"><v>43000.75</v></c></row>'
    )
    path = tmp_path_factory.mktemp("compat") / "epoch1904.xlsx"
    return build(path, {"S": sheet(rows, dimension="A1:C1")}, date1904=True)


# ---- the flag ----------------------------------------------------------------------


def test_native_api_unchanged():
    path = DATA / "base.xlsx"
    native = load_workbook(path)
    assert type(native) is CalamineWorkbook
    rows = list(native.stream_sheet_by_index(0))
    assert rows == list(CalamineWorkbook.from_path(path).stream_sheet_by_index(0))
    assert "" in rows[0] or all(v != "" for v in rows[0])  # native empties stay ""
    assert type(load_workbook(path, compatibility=None)) is CalamineWorkbook


def test_options_are_validated(tmp_path):
    path = DATA / "base.xlsx"
    with pytest.raises(TypeError, match="compatibility"):
        load_workbook(path, read_only=True)
    with pytest.raises(TypeError, match="compatibility"):
        load_workbook(path, data_only=False)
    with pytest.raises(ValueError, match="unknown compatibility"):
        load_workbook(path, compatibility="xlrd")
    with pytest.raises(ValueError, match="keep_vba"):
        load_workbook(path, compatibility="openpyxl", keep_vba=True)
    with pytest.raises(ValueError, match="keep_links"):
        load_workbook(path, compatibility="openpyxl", keep_links=False)
    with pytest.raises(CompatibilityNotSupported, match="rich_text"):
        load_workbook(path, compatibility="openpyxl", rich_text=True)
    with pytest.raises(ValueError, match="load_tables"):
        load_workbook(path, True, compatibility="openpyxl")
    # Defaults are accepted.
    load_workbook(
        path, compatibility="openpyxl", keep_vba=False, keep_links=True, rich_text=False
    ).close()


@pytest.mark.parametrize("name", ["base.xls", "base.xlsb", "base.ods"])
def test_other_formats_rejected(name):
    with pytest.raises(CompatibilityNotSupported, match="xlsx"):
        load_workbook(DATA / name, compatibility="openpyxl")
    with pytest.raises(CompatibilityNotSupported):
        load_workbook((DATA / name).open("rb"), compatibility="openpyxl")


def test_drop_in_module(styled_xlsx):
    from python_calamine_plus import openpyxl as compat

    wb = compat.load_workbook(styled_xlsx, data_only=True)
    assert wb["Data"]["A4"].value == 1
    with pytest.raises(ValueError):
        compat.load_workbook(styled_xlsx, keep_vba=True)
    from python_calamine_plus.openpyxl.utils import get_column_letter
    from python_calamine_plus.openpyxl.utils.cell import coordinate_to_tuple
    from python_calamine_plus.openpyxl.utils.datetime import from_excel

    assert get_column_letter(28) == "AB"
    assert coordinate_to_tuple("C7") == (7, 3)
    assert from_excel(1) == datetime.datetime(1900, 1, 1)


# ---- parity ------------------------------------------------------------------------


@pytest.mark.parametrize("read_only,data_only", MODES)
def test_parity_styled(styled_xlsx, read_only, data_only):
    assert_parity(styled_xlsx, read_only, data_only)


@pytest.mark.parametrize("read_only,data_only", MODES)
def test_parity_raw_values(raw_xlsx, read_only, data_only):
    ox = openpyxl.load_workbook(raw_xlsx, read_only=read_only, data_only=data_only)
    cp = load_workbook(
        raw_xlsx, compatibility="openpyxl", read_only=read_only, data_only=data_only
    )
    for title in ("Values", "NoDim"):
        o = [[_value(v) for v in r] for r in ox[title].iter_rows(values_only=True)]
        c = [[_value(v) for v in r] for r in cp[title].iter_rows(values_only=True)]
        assert c == o, title
        o = [[_cell(x) for x in r] for r in ox[title].iter_rows()]
        c = [[_cell(x) for x in r] for r in cp[title].iter_rows()]
        assert c == o, title


@pytest.mark.parametrize("read_only,data_only", MODES)
@pytest.mark.parametrize(
    "name", ["base.xlsx", "any_sheets.xlsx", "issue139.xlsx", "table-multiple.xlsx"]
)
def test_parity_test_files(name, read_only, data_only):
    assert_parity(DATA / name, read_only, data_only)


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize(
    "bounds",
    [
        dict(min_row=2, max_row=6, min_col=2, max_col=5),
        dict(min_row=3),
        dict(max_col=2),
        dict(min_col=4, max_row=15),
        dict(min_row=1, max_row=40, max_col=12),
    ],
)
def test_parity_bounded_iteration(styled_xlsx, read_only, bounds):
    assert_parity(styled_xlsx, read_only, False, **bounds)


@pytest.mark.parametrize("read_only,data_only", MODES)
def test_parity_merged_hyperlinks(raw_xlsx, read_only, data_only):
    ox = openpyxl.load_workbook(raw_xlsx, read_only=read_only, data_only=data_only)[
        "Merged"
    ]
    cp = load_workbook(
        raw_xlsx, compatibility="openpyxl", read_only=read_only, data_only=data_only
    )["Merged"]
    assert [[_cell(c) for c in r] for r in cp.iter_rows()] == [
        [_cell(c) for c in r] for r in ox.iter_rows()
    ]
    if not read_only:
        assert cp["A1"].hyperlink.location == "Values!A1"
        assert cp["B2"].hyperlink is None and cp["B1"].value is None
        assert cp["C3"].value == "Values!B1"  # an empty linked cell shows the link
        assert cp.dimensions == ox.dimensions == "A1:C3"


def test_unstored_cells_use_openpyxl_zero_style(tmp_path):
    """Cells the file does not store (and merged placeholders) get openpyxl's
    all-zero style array: default alignment / protection, not cellXfs[0]'s."""
    styles_xf0 = (
        '<row r="1"><c r="A1" s="1"><v>1</v></c></row>'
        '<row r="3"><c r="C3"><v>2</v></c></row>'
    )
    path = build(
        tmp_path / "xf0.xlsx",
        {
            "S": sheet(
                styles_xf0,
                dimension="A1:C3",
                extra='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
            )
        },
    )
    # cellXfs[0] with an explicit alignment and protection.
    import zipfile

    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    files["xl/styles.xml"] = files["xl/styles.xml"].replace(
        b'<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>',
        b'<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1">'
        b'<alignment horizontal="general" vertical="bottom"/><protection locked="0"/></xf>',
        1,
    )
    with zipfile.ZipFile(path, "w") as z:
        for n, data in files.items():
            z.writestr(n, data)
    ox = openpyxl.load_workbook(path)["S"]
    cp = load_workbook(path, compatibility="openpyxl")["S"]
    for coord in ("A1", "B1", "B2", "A3", "C3", "D9"):
        assert _cell(cp[coord]) == _cell(ox[coord]), coord
    assert cp["C3"].alignment.vertical == "bottom"  # stored with style 0
    assert cp["A3"].alignment.vertical is None  # not stored: openpyxl's default
    assert cp["A3"].protection.locked is True and cp["A3"].style_id == 0


def test_parity_epoch_1904(epoch_1904_xlsx):
    ox = openpyxl.load_workbook(epoch_1904_xlsx)
    cp = load_workbook(epoch_1904_xlsx, compatibility="openpyxl")
    assert cp.epoch == ox.epoch == datetime.datetime(1904, 1, 1)
    assert list(cp["S"].values) == list(ox["S"].values)


def test_values(raw_xlsx):
    ws = load_workbook(raw_xlsx, compatibility="openpyxl")["Values"]
    row1 = next(ws.iter_rows(max_row=1, values_only=True))
    assert row1[:3] == (5, 5.0, 1000.0)
    assert [type(v) for v in row1[:3]] == [int, float, float]
    assert row1[3] == datetime.datetime(
        2024, 2, 29
    )  # dates are datetimes, like openpyxl
    assert row1[4] == datetime.timedelta(hours=30)
    assert row1[5] == datetime.time(12)
    assert row1[6] == 12  # "days" 0 is not a date format
    assert row1[8] == datetime.datetime(
        1900, 2, 28
    )  # serial 59 (before the 1900 leap bug)
    assert row1[9] == datetime.datetime(1900, 3, 1)
    a2, b2, c2, d2, e2, f2, g2, h2, i2, j2 = ws[2][:10]
    assert (a2.value, b2.value, c2.value, d2.value) == ("shared", "", "inline", "")
    assert (e2.data_type, e2.value) == ("f", '="x"&"y"')
    assert (f2.value, g2.value, g2.data_type) == (False, "#DIV/0!", "e")
    assert h2.value == datetime.datetime(2024, 1, 31, 10, 30)
    assert i2.value == datetime.date(2024, 1, 31)
    assert j2.value is None and j2.font.b is True  # stored, styled, empty


def test_formulas_and_cached_values(raw_xlsx):
    formulas = load_workbook(raw_xlsx, compatibility="openpyxl")["Values"]
    cached = load_workbook(raw_xlsx, compatibility="openpyxl", data_only=True)["Values"]
    assert [c.value for c in formulas[4][:3]] == ["=A1+1", "=B1+1", "=C1+1"]
    assert [c.data_type for c in formulas[4][:3]] == ["f"] * 3
    assert [c.value for c in cached[4][:3]] == [6, 6, 1001]
    d4 = formulas["D4"].value
    assert (type(d4).__name__, d4.ref, d4.text, d4.t) == (
        "ArrayFormula",
        "D4:D5",
        "=A1:A2*2",
        "array",
    )
    assert cached["D4"].value == 10 and cached["D5"].value == 4
    table = formulas["B6"].value
    assert (table.t, table.ref, table.dt2D, table.dtr, table.r1, table.ca) == (
        "dataTable",
        "B6:B7",
        "0",
        "1",
        "A1",
        False,
    )
    # A formula without a saved result has no cached value (nothing is recalculated).
    assert formulas["E4"].value == "=A1/0"
    assert cached["E4"].value is None
    assert (cached["F4"].value, cached["F4"].data_type) == ("#DIV/0!", "e")
    assert formulas["E4"].data_type == "f" and cached["E4"].data_type == "n"
    # An empty cached string is no value (openpyxl keeps the "str" type).
    assert (cached["G4"].value, cached["G4"].data_type) == (None, "str")


# ---- bounds, padding, random access ------------------------------------------------


def test_declared_vs_observed_bounds(raw_xlsx, styled_xlsx):
    ro = load_workbook(raw_xlsx, compatibility="openpyxl", read_only=True)
    normal = load_workbook(raw_xlsx, compatibility="openpyxl")
    # Read-only: the declared <dimension>; normal: the stored cells.
    assert (ro["Values"].max_row, ro["Values"].max_column) == (12, 11)
    assert (normal["Values"].max_row, normal["Values"].max_column) == (8, 11)
    assert ro["NoDim"].max_row is None
    with pytest.raises(ValueError, match="unsized"):
        ro["NoDim"].calculate_dimension()
    assert ro["NoDim"].calculate_dimension(force=True) == "A1:D4"
    assert normal["NoDim"].dimensions == "B2:D4"
    # Styled empty cells, merged ranges, hyperlinks and comments count (normal mode).
    ws = load_workbook(styled_xlsx, compatibility="openpyxl")["Data"]
    assert ws.dimensions == "A1:J12"
    assert (
        load_workbook(styled_xlsx, compatibility="openpyxl")["Empty"].dimensions
        == "A1:A1"
    )


def test_rows_are_rectangular(raw_xlsx):
    for read_only in (False, True):
        ws = load_workbook(raw_xlsx, compatibility="openpyxl", read_only=read_only)[
            "Values"
        ]
        rows = list(
            ws.iter_rows(min_row=2, max_row=9, min_col=2, max_col=5, values_only=True)
        )
        assert len(rows) == 8 and {len(r) for r in rows} == {4}
        assert rows[1] == (None,) * 4  # row 3 is not stored
        cells = list(ws.iter_rows(min_row=2, max_row=9, min_col=2, max_col=5))
        assert [[c.value for c in r] for r in cells] == [list(r) for r in rows]


def test_random_access(styled_xlsx):
    ws = load_workbook(styled_xlsx, compatibility="openpyxl")["Data"]
    assert (
        ws["A4"].value == 1
        and ws.cell(4, 2).value == 2.5
        and ws.cell(row=4, column=1).row == 4
    )
    assert ws["Z99"].value is None and ws["Z99"].coordinate == "Z99"
    assert ws.max_column == 10  # access does not create cells
    assert [c.value for c in ws["A4:B4"][0]] == [1, 2.5]
    assert [c.coordinate for c in ws["A"]][:2] == ["A1", "A2"]
    assert len(ws[4]) == ws.max_column
    assert [len(r) for r in ws[4:5]] == [10, 10]
    assert [len(c) for c in ws["A:B"]] == [12, 12]
    assert ws["A4"].offset(0, 1).value == 2.5
    assert type(ws["B1"]).__name__ == "MergedCell" and ws["B1"].value is None
    with pytest.raises(ValueError):
        ws.cell(0, 1)
    with pytest.raises(AttributeError):
        ws.cell(1, 1, value=3)
    with pytest.raises(AttributeError):
        ws["A4"].value = 3
    # Backwards access restarts the reader; results are unchanged.
    assert ws["J12"].comment.text == "far away" and ws["A4"].value == 1


def test_sparse_sheet_is_not_materialized(raw_xlsx):
    for read_only in (False, True):
        ws = load_workbook(raw_xlsx, compatibility="openpyxl", read_only=read_only)[
            "Sparse"
        ]
        assert (ws.max_row, ws.max_column) == (1_048_576, 16_384)
        assert ws["XFD1048576"].value == 2
        assert list(ws.iter_rows(max_row=2, max_col=2, values_only=True)) == [
            (1, None),
            (None, None),
        ]


# ---- metadata ----------------------------------------------------------------------


def test_metadata(styled_xlsx):
    wb = load_workbook(styled_xlsx, compatibility="openpyxl")
    ws = wb["Data"]
    assert sorted(str(r) for r in ws.merged_cells.ranges) == ["A1:C2", "E1:F3"]
    assert str(ws.merged_cells) == "A1:C2 E1:F3"
    assert "B2" in ws.merged_cells and "D4" not in ws.merged_cells
    assert ws.merged_cells.sorted()[0].start_cell.value == "Header"
    assert ws.print_area == "'Data'!$A$1:$C$10,'Data'!$E$1:$F$2"
    assert ws.print_titles == "'Data'!$1:$2,'Data'!$A:$B"
    assert ws.row_dimensions[3].height == 30 and ws.row_dimensions[20].hidden
    assert ws.row_dimensions[99].height is None  # default, not stored
    assert ws.column_dimensions["B"].width == 25
    assert ws.column_dimensions["D"].range == "D:F"
    assert ws.freeze_panes == "B2"
    assert ws.sheet_properties.tabColor.rgb == "FF123456"
    assert wb["Hidden"].sheet_state == "hidden"
    assert list(ws.defined_names) == ["local_name"]
    assert wb.defined_names["global_name"].value == "Data!$B$4:$C$4"
    assert list(wb.defined_names["global_name"].destinations) == [("Data", "$B$4:$C$4")]
    assert ws["A8"].comment.author == "Ann"
    assert ws["B7"].value == "https://example.com/empty" and ws["B7"].data_type == "s"
    assert ws["A7"].hyperlink.target == "https://example.com/a?x=1&y=2"
    assert ws["A1"].hyperlink is None
    with pytest.raises(AttributeError, match="compatibility"):
        ws.conditional_formatting
    with pytest.raises(AttributeError):
        ws.no_such_attribute


def test_merged_borders_match_openpyxl(styled_xlsx):
    ox = openpyxl.load_workbook(styled_xlsx)["Data"]
    cp = load_workbook(styled_xlsx, compatibility="openpyxl")["Data"]
    for coord in ("A1", "B1", "C1", "A2", "B2", "C2", "E1", "F1", "E3", "F3", "F2"):
        assert type(cp[coord]).__name__ == type(ox[coord]).__name__, coord
        assert _style(cp[coord]) == _style(ox[coord]), coord


def test_tables():
    ws = load_workbook(DATA / "table-multiple.xlsx", compatibility="openpyxl")["Sheet1"]
    ox = openpyxl.load_workbook(DATA / "table-multiple.xlsx")["Sheet1"]
    assert sorted(ws.tables.items()) == sorted(ox.tables.items())
    for name in ox.tables:
        assert ws.tables[name].column_names == ox.tables[name].column_names
        assert ws.tables[name].displayName == ox.tables[name].displayName


# ---- streaming, iterators, lifecycle -----------------------------------------------


def test_iterators_are_independent(styled_xlsx):
    for read_only in (False, True):
        ws = load_workbook(styled_xlsx, compatibility="openpyxl", read_only=read_only)[
            "Data"
        ]
        expected = list(ws.iter_rows(values_only=True))
        a = ws.iter_rows(values_only=True)
        b = ws.iter_rows(values_only=True)
        out_a, out_b = [], []
        for ra, rb in itertools.zip_longest(a, b):
            out_a.append(ra)
            out_b.append(rb)
            ws.cell(1, 1)  # random access in between
            list(ws.iter_rows(max_row=1))  # and a third, complete iterator
        assert out_a == out_b == expected


def test_read_only_values_only_reads_no_metadata(raw_xlsx, styled_xlsx):
    wb = load_workbook(styled_xlsx, compatibility="openpyxl", read_only=True)
    ws = wb["Data"]
    list(ws.iter_rows(values_only=True))
    list(ws.iter_rows())
    assert ws._metadata_loaded is False
    ws.merged_cells  # metadata is read on demand
    assert ws._metadata_loaded is True


def test_close(styled_xlsx):
    for read_only in (False, True):
        wb = load_workbook(styled_xlsx, compatibility="openpyxl", read_only=read_only)
        ws = wb["Data"]
        rows = ws.iter_rows(values_only=True)
        first = next(rows)
        cell = ws["A4"]
        wb.close()
        assert first[0] == "Header" and cell.value == 1 and cell.font.b is False
        with pytest.raises(WorkbookClosed):
            next(rows)
        with pytest.raises(WorkbookClosed):
            ws["A5"]
        with pytest.raises(WorkbookClosed):
            list(ws.iter_rows())
        wb.close()  # closing twice is fine
    with load_workbook(styled_xlsx, compatibility="openpyxl") as wb:
        assert wb["Data"]["A4"].value == 1
    with pytest.raises(WorkbookClosed):
        wb["Data"]["A5"]


def test_close_releases_the_file(styled_xlsx, tmp_path):
    import gc
    import shutil

    for read_only in (False, True):
        path = tmp_path / f"copy-{read_only}.xlsx"
        shutil.copy(styled_xlsx, path)
        wb = load_workbook(path, compatibility="openpyxl", read_only=read_only)
        ws = wb["Data"]
        assert ws["A4"].value == 1  # opens the random-access reader
        list(ws.iter_rows(values_only=True))  # an exhausted iterator
        rows = ws.iter_rows()
        next(rows)  # an unfinished iterator holds its own reader...
        wb.close()
        with pytest.raises(WorkbookClosed):
            next(rows)  # ...until it is used after close
        del rows
        gc.collect()
        path.unlink()  # fails on Windows while any handle is open


def test_filelike(styled_xlsx):
    with styled_xlsx.open("rb") as f:
        wb = load_workbook(f, compatibility="openpyxl", data_only=True)
    from_path = load_workbook(styled_xlsx, compatibility="openpyxl", data_only=True)
    assert list(wb["Data"].values) == list(from_path["Data"].values)


def test_workbook_api(styled_xlsx):
    wb = load_workbook(styled_xlsx, compatibility="openpyxl")
    assert wb.sheetnames == ["Data", "Hidden", "Empty"]
    assert [ws.title for ws in wb] == wb.sheetnames
    assert "Data" in wb and "Nope" not in wb
    assert wb.index(wb["Hidden"]) == 1
    assert wb["Data"] is wb["Data"] and wb["Data"].parent is wb
    with pytest.raises(KeyError, match="does not exist"):
        wb["Nope"]
    assert wb.read_only is False and wb.data_only is False
    assert wb.epoch == datetime.datetime(1899, 12, 30)
    assert "Normal" in wb.named_styles
    assert wb.properties.creator == "openpyxl"
    assert (
        list(load_workbook(styled_xlsx, compatibility="openpyxl")["Empty"].iter_rows())
        == []
    )
    with pytest.raises(TypeError, match="read-only"):
        wb.save("x.xlsx")
    assert repr(wb["Data"]) == '<Worksheet "Data">'
    assert repr(wb["Data"]["A4"]) == "<Cell 'Data'.A4>"


def test_style_objects_are_shared(styled_xlsx):
    """Style objects (and their colors / sides) are built once per cell format."""
    for read_only in (False, True):
        ws = load_workbook(styled_xlsx, compatibility="openpyxl", read_only=read_only)[
            "Data"
        ]
        a, b = ws["A4"], ws["B4"]  # same (default) format
        assert a.font is b.font and a.fill is b.fill
        assert a.fill.fgColor is a.fill.fgColor and a.border.left is a.border.left
        assert a.number_format is b.number_format
        header = ws["A1"]
        assert header.font.color is header.font.color
        assert header.font.color.rgb == "FFFF0000"
