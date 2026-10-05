"""Opt-in extensions of compatibility="openpyxl": formula_and_value, and read-only
comments / hyperlinks / merged cells (like the native rich stream)."""

from __future__ import annotations

import pytest
from python_calamine_plus import CalamineWorkbook, load_workbook

from .xlsx_builder import build, sheet

openpyxl = pytest.importorskip("openpyxl")

MODES = [(False, False), (False, True), (True, False), (True, True)]


@pytest.fixture(scope="module")
def formulas_xlsx(tmp_path_factory):
    rows = (
        '<row r="1"><c r="A1"><v>5</v></c><c r="B1"><f>A1*2</f><v>10</v></c>'
        '<c r="C1" t="str"><f>"a"&amp;"b"</f><v>ab</v></c><c r="D1"><f>A1/0</f></c>'
        '<c r="E1" t="s"><v>0</v></c></row>'
        '<row r="2"><c r="A2"><f t="array" ref="A2:A3">A1:A2*2</f><v>10</v></c></row>'
        '<row r="3"><c r="A3"><v>7</v></c></row>'
    )
    return build(
        tmp_path_factory.mktemp("ext") / "formulas.xlsx",
        {"F": sheet(rows, dimension="A1:E3")},
        shared_strings=["text"],
    )


@pytest.fixture(scope="module")
def notes_xlsx(tmp_path_factory):
    from openpyxl.comments import Comment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "title"
    ws.merge_cells("A1:C1")
    ws["A2"] = "noted"
    ws["A2"].comment = Comment("on a value", "Ann")
    ws["D2"].comment = Comment("on an empty cell", "Bo")
    ws["B3"] = "linked"
    ws["B3"].hyperlink = "https://example.com/?a=1&b=2"
    ws["A4"] = 1
    path = tmp_path_factory.mktemp("ext") / "notes.xlsx"
    wb.save(path)
    return path


# ---- formula_and_value -------------------------------------------------------------


@pytest.mark.parametrize("read_only,data_only", MODES)
def test_formula_and_value(formulas_xlsx, read_only, data_only):
    wb = load_workbook(
        formulas_xlsx,
        compatibility="openpyxl",
        read_only=read_only,
        data_only=data_only,
        formula_and_value=True,
    )
    plain = load_workbook(
        formulas_xlsx,
        compatibility="openpyxl",
        read_only=read_only,
        data_only=data_only,
    )
    rows = list(wb["F"].iter_rows())
    a1, b1, c1, d1, e1 = rows[0]
    # `value` / `data_type` are unchanged by the flag.
    expected = [[(c.value, c.data_type) for c in r] for r in plain["F"].iter_rows()]
    got = [[(c.value, c.data_type) for c in r] for r in rows]
    assert [[v for v in r if not hasattr(v[0], "text")] for r in got] == [
        [v for v in r if not hasattr(v[0], "text")] for r in expected
    ]
    # Formula and saved result from the same pass.
    assert (b1.formula, b1.cached_value) == ("=A1*2", 10)
    assert (c1.formula, c1.cached_value) == ('="a"&"b"', "ab")
    assert (d1.formula, d1.cached_value) == ("=A1/0", None)  # never calculated
    assert (a1.formula, a1.cached_value) == (None, 5)
    assert (e1.formula, e1.cached_value) == (None, "text")
    array = rows[1][0].formula
    assert (array.ref, array.text, rows[1][0].cached_value) == ("A2:A3", "=A1:A2*2", 10)
    # The same through random access.
    assert wb["F"]["B1"].formula == "=A1*2" and wb["F"]["B1"].cached_value == 10


def test_formula_and_value_matches_openpyxl_views(formulas_xlsx):
    both = load_workbook(
        formulas_xlsx, compatibility="openpyxl", formula_and_value=True
    )["F"]
    formulas = openpyxl.load_workbook(formulas_xlsx)["F"]
    values = openpyxl.load_workbook(formulas_xlsx, data_only=True)["F"]
    for row in both.iter_rows():
        for c in row:
            f = formulas[c.coordinate]
            expected_formula = f.value if f.data_type == "f" else None
            if hasattr(expected_formula, "text"):
                assert (c.formula.ref, c.formula.text) == (
                    expected_formula.ref,
                    expected_formula.text,
                )
            else:
                assert c.formula == expected_formula, c.coordinate
            assert c.cached_value == values[c.coordinate].value, c.coordinate


def test_formula_and_value_is_opt_in(formulas_xlsx):
    for read_only in (False, True):
        cell = load_workbook(
            formulas_xlsx, compatibility="openpyxl", read_only=read_only
        )["F"]["B1"]
        with pytest.raises(AttributeError, match="formula_and_value=True"):
            cell.formula
        with pytest.raises(AttributeError, match="formula_and_value=True"):
            cell.cached_value


def test_formula_and_value_drop_in(formulas_xlsx):
    from python_calamine_plus import openpyxl as compat

    ws = compat.load_workbook(formulas_xlsx, read_only=True, formula_and_value=True)[
        "F"
    ]
    assert ws["B1"].formula == "=A1*2" and ws["B1"].cached_value == 10


# ---- read-only comments / hyperlinks / merged cells --------------------------------


def _ro(path, **flags):
    return load_workbook(path, compatibility="openpyxl", read_only=True, **flags)["S"]


def test_read_only_default_follows_openpyxl(notes_xlsx):
    ws = _ro(notes_xlsx)
    cell = ws["A2"]
    for attr, flag in (
        ("comment", "read_comments"),
        ("hyperlink", "read_hyperlinks"),
        ("merged_range", "read_merged_cells"),
    ):
        with pytest.raises(AttributeError, match=flag):
            getattr(cell, attr)
    # Without the flags, cells are exactly openpyxl's read-only cells.
    ox = openpyxl.load_workbook(notes_xlsx, read_only=True)["S"]
    assert [[type(c).__name__ for c in r] for r in ws.iter_rows()] == [
        [type(c).__name__ for c in r] for r in ox.iter_rows()
    ]


def test_read_only_opt_in(notes_xlsx):
    ws = _ro(
        notes_xlsx, read_comments=True, read_hyperlinks=True, read_merged_cells=True
    )
    rows = list(ws.iter_rows())
    a2 = rows[1][0]
    assert (a2.comment.text, a2.comment.author) == ("on a value", "Ann")
    d2 = rows[1][3]  # no stored cell, but a comment: a cell like the native stream
    assert type(d2).__name__ == "ReadOnlyCell" and d2.value is None
    assert d2.comment.text == "on an empty cell"
    b3 = rows[2][1]
    assert (
        b3.hyperlink.target == "https://example.com/?a=1&b=2" and b3.value == "linked"
    )
    assert [str(c.merged_range) for c in rows[0][:3]] == ["A1:C1"] * 3
    assert rows[0][0].hyperlink is None and rows[3][0].merged_range is None
    # Values-only rows are unchanged.
    assert list(ws.iter_rows(values_only=True)) == list(
        _ro(notes_xlsx).iter_rows(values_only=True)
    )
    assert ws["D2"].comment.author == "Bo"


def test_read_only_flags_match_native_rich_stream(notes_xlsx):
    ws = _ro(
        notes_xlsx, read_comments=True, read_hyperlinks=True, read_merged_cells=True
    )
    native = {
        (c.row + 1, c.column + 1): c
        for row in CalamineWorkbook.from_path(notes_xlsx).stream_sheet_by_name(
            "S", rich=True
        )
        for c in row
    }
    for row in ws.iter_rows():
        for c in row:
            if type(c).__name__ == "EmptyCell":
                continue
            n = native[(c.row, c.column)]
            assert (c.comment.text if c.comment else None) == (
                n.comment["text"] if n.comment else None
            )
            assert (c.hyperlink.target if c.hyperlink else None) == (
                n.hyperlink["target"] if n.hyperlink else None
            )
            expected = n.merged_range
            got = c.merged_range
            assert (
                None
                if got is None
                else (
                    (got.min_row - 1, got.min_col - 1),
                    (got.max_row - 1, got.max_col - 1),
                )
            ) == expected


def test_metadata_flags_are_validated(notes_xlsx):
    with pytest.raises(ValueError, match="read_only=True"):
        load_workbook(notes_xlsx, compatibility="openpyxl", read_comments=False)
    with pytest.raises(TypeError, match="compatibility"):
        load_workbook(notes_xlsx, read_comments=True)
    # Normal mode always has them (True is accepted).
    ws = load_workbook(notes_xlsx, compatibility="openpyxl", read_comments=True)["S"]
    assert ws["A2"].comment.text == "on a value"


# ---- bugs found by the corpus comparison -------------------------------------------


def test_shared_formula_relative_to_its_master_cell(tmp_path):
    """The master formula is relative to the cell holding it, which need not be the
    first cell of the shared range (here B1 in A1:D1)."""
    rows = (
        '<row r="1"><c r="B1"><f t="shared" ref="A1:D1" si="0">B2*2</f><v>0</v></c>'
        '<c r="C1"><f t="shared" si="0"/><v>0</v></c>'
        '<c r="D1"><f t="shared" si="0"/><v>0</v></c></row>'
    )
    path = build(tmp_path / "shared.xlsx", {"S": sheet(rows, dimension="B1:D1")})
    expected = [c.value for c in openpyxl.load_workbook(path)["S"][1][1:]]
    assert expected == ["=B2*2", "=C2*2", "=D2*2"]
    ws = load_workbook(path, compatibility="openpyxl")["S"]
    assert [c.value for c in ws[1][1:]] == expected
    native = next(CalamineWorkbook.from_path(path).stream_sheet_by_name("S", rich=True))
    assert [c.formula for c in native[1:]] == expected


def test_hyperlink_attributes_are_entity_decoded(tmp_path):
    rows = '<row r="1"><c r="A1"><v>1</v></c></row>'
    links = (
        '<hyperlinks><hyperlink ref="A1" location="&apos;My Sheet&apos;!A1" '
        'display="a &amp; b" tooltip="&lt;tip&gt;"/></hyperlinks>'
    )
    path = build(
        tmp_path / "links.xlsx", {"S": sheet(rows, dimension="A1", extra=links)}
    )
    h = load_workbook(path, compatibility="openpyxl")["S"]["A1"].hyperlink
    o = openpyxl.load_workbook(path)["S"]["A1"].hyperlink
    assert (h.location, h.display, h.tooltip) == (o.location, o.display, o.tooltip)
    assert h.location == "'My Sheet'!A1"
    native = next(
        CalamineWorkbook.from_path(path).stream_sheet_by_name("S", rich=True)
    )[0]
    assert native.hyperlink["location"] == "'My Sheet'!A1"
    assert (native.hyperlink["display"], native.hyperlink["tooltip"]) == (
        "a & b",
        "<tip>",
    )
