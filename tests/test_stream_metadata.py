from datetime import datetime
from io import BytesIO

import pytest
from python_calamine_plus import CalamineWorkbook

openpyxl = pytest.importorskip("openpyxl")


@pytest.fixture(scope="module")
def meta_xlsx(tmp_path_factory):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.formatting.rule import (
        CellIsRule,
        ColorScaleRule,
        DataBarRule,
        IconSetRule,
    )
    from openpyxl.styles import Font, PatternFill
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.hyperlink import Hyperlink

    wb = openpyxl.Workbook()
    wb.properties.creator = "Amir"
    wb.properties.title = "Fork metadata"
    wb.properties.created = datetime(2024, 1, 31, 10, 30, 0)
    ws = wb.active
    ws.title = "Meta"
    other = wb.create_sheet("Other")
    other["A1"] = "x"

    for r in range(1, 21):
        ws.cell(row=r, column=1, value=r)
        ws.cell(row=r, column=2, value=r * 1.5)
    ws["C1"] = "#DIV/0!"  # error cell
    ws["C2"] = "#N/A"
    ws["C3"] = 42
    ws["C4"] = 42.0
    ws["C5"] = -7
    ws["D1"] = "Docs"
    ws["D1"].hyperlink = "https://example.com/docs"
    ws["D2"] = "Jump"
    ws["D2"].hyperlink = Hyperlink(ref="D2", location="Other!A1", tooltip="Go")
    ws["D3"] = "Outside"
    ws["D3"].hyperlink = "#Other!A1"  # openpyxl stores this as an external target
    ws["E1"] = CellRichText(
        "plain ",
        TextBlock(InlineFont(b=True, color="FFFF0000"), "bold red"),
        " tail",
    )
    ws["E2"] = "not rich"

    ws.conditional_formatting.add(
        "A1:A20",
        CellIsRule(
            operator="greaterThan",
            formula=["10"],
            font=Font(bold=True, color="FF9C0006"),
            fill=PatternFill(bgColor="FFFFC7CE"),
        ),
    )
    ws.conditional_formatting.add(
        "B1:B20",
        ColorScaleRule(
            start_type="min",
            start_color="FFF8696B",
            end_type="max",
            end_color="FF63BE7B",
        ),
    )
    ws.conditional_formatting.add(
        "B1:B20", DataBarRule(start_type="min", end_type="max", color="FF638EC6")
    )
    ws.conditional_formatting.add(
        "A1:A20", IconSetRule("3TrafficLights1", "percent", [0, 33, 67])
    )

    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv.prompt = "Pick one"
    dv.promptTitle = "Answer"
    ws.add_data_validation(dv)
    dv.add("F1:F10")
    whole = DataValidation(
        type="whole", operator="between", formula1="1", formula2="10"
    )
    whole.error = "1 to 10"
    ws.add_data_validation(whole)
    whole.add("G1")

    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].hidden = True
    ws.column_dimensions.group("E", "F", outline_level=1)
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].hidden = True
    ws.row_dimensions[4].outlineLevel = 2
    ws.row_dimensions[25].hidden = True  # empty row after the data: never streamed
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:B20"
    ws.protection.sheet = True
    ws.protection.sort = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_margins.left = 0.5
    ws.print_options.gridLines = True
    ws.sheet_properties.tabColor = "FF1072BA"
    ws.sheet_view.showGridLines = False

    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Meta!$B$2")
    ws.defined_names["LocalName"] = DefinedName("LocalName", attr_text="Meta!$A$1:$A$5")

    # A hidden empty row inside the data.
    ws2 = wb.create_sheet("Gaps")
    ws2["A1"] = 1
    ws2.row_dimensions[2].hidden = True
    ws2["A3"] = 3

    path = tmp_path_factory.mktemp("meta") / "meta.xlsx"
    wb.save(path)
    return path


def _rich(path, sheet="Meta", **kw):
    return CalamineWorkbook.from_path(path).stream_sheet_by_name(sheet, rich=True, **kw)


# ---- values -------------------------------------------------------------------


def test_error_values_are_strings(meta_xlsx):
    wb = CalamineWorkbook.from_path(meta_xlsx)
    plain = list(wb.stream_sheet_by_name("Meta"))
    assert plain[0][2] == "#DIV/0!"
    assert plain[1][2] == "#N/A"
    rich = list(wb.stream_sheet_by_name("Meta", rich=True))
    assert rich[0][2].value == "#DIV/0!"
    full = wb.get_sheet_by_name("Meta").to_python()
    assert full[0][2] == "#DIV/0!"


def test_integers_are_int(meta_xlsx):
    rows = list(CalamineWorkbook.from_path(meta_xlsx).stream_sheet_by_name("Meta"))
    assert (
        isinstance(rows[0][0], int)
        and not isinstance(rows[0][0], bool)
        and rows[0][0] == 1
    )
    assert (
        isinstance(rows[2][2], int)
        and not isinstance(rows[2][2], bool)
        and rows[2][2] == 42
    )
    assert (
        isinstance(rows[4][2], int)
        and not isinstance(rows[4][2], bool)
        and rows[4][2] == -7
    )
    assert isinstance(rows[0][1], float) and rows[0][1] == 1.5
    ox = openpyxl.load_workbook(meta_xlsx, read_only=True)["Meta"]
    for ours, theirs in zip(rows, ox.iter_rows(values_only=True)):
        for a, b in zip(ours, theirs):
            if isinstance(b, (int, float)) and not isinstance(b, bool):
                assert type(a) is type(b), (a, b)


# ---- workbook level -------------------------------------------------------------


def test_defined_names(meta_xlsx):
    names = {n["name"]: n for n in CalamineWorkbook.from_path(meta_xlsx).defined_names}
    assert names["TaxRate"] == {
        "name": "TaxRate",
        "value": "Meta!$B$2",
        "sheet": None,
        "hidden": False,
        "comment": None,
    }
    assert names["LocalName"]["sheet"] == "Meta"
    assert names["LocalName"]["value"] == "Meta!$A$1:$A$5"
    # The auto filter creates a hidden, sheet-scoped _FilterDatabase name.
    assert names["_xlnm._FilterDatabase"]["hidden"] is True


def test_properties(meta_xlsx):
    props = CalamineWorkbook.from_path(meta_xlsx).properties
    assert props["creator"] == "Amir"
    assert props["title"] == "Fork metadata"
    assert props["created"] == datetime(2024, 1, 31, 10, 30, 0)
    assert isinstance(props["modified"], datetime)


def test_workbook_info_from_filelike(meta_xlsx):
    wb = CalamineWorkbook.from_filelike(BytesIO(meta_xlsx.read_bytes()))
    assert wb.properties["creator"] == "Amir"
    assert any(n["name"] == "TaxRate" for n in wb.defined_names)


# ---- per cell -------------------------------------------------------------------


def test_hyperlinks(meta_xlsx):
    stream = _rich(meta_xlsx)
    rows = list(stream)
    d1, d2 = rows[0][3], rows[1][3]
    assert d1.hyperlink["target"] == "https://example.com/docs"
    assert d1.hyperlink["range"] == ((0, 3), (0, 3))
    assert d2.hyperlink["location"] == "Other!A1"
    assert d2.hyperlink["target"] is None
    assert d2.hyperlink["tooltip"] == "Go"
    assert rows[2][3].hyperlink["target"] == "#Other!A1"
    assert rows[0][0].hyperlink is None
    assert len(stream.hyperlinks) == 3


def test_rich_text(meta_xlsx):
    rows = list(_rich(meta_xlsx))
    e1 = rows[0][4]
    assert e1.value == "plain bold red tail"
    runs = e1.rich_text
    assert [r["text"] for r in runs] == ["plain ", "bold red", " tail"]
    assert runs[0]["font"] is None
    assert runs[1]["font"]["bold"] is True
    assert runs[1]["font"]["color"] == {"rgb": "FFFF0000"}
    assert rows[1][4].value == "not rich" and rows[1][4].rich_text is None
    assert rows[0][0].rich_text is None
    # Same flattened text as the plain stream.
    plain = list(CalamineWorkbook.from_path(meta_xlsx).stream_sheet_by_name("Meta"))
    assert plain[0][4] == "plain bold red tail"


# ---- per row --------------------------------------------------------------------


def test_row_info(meta_xlsx):
    stream = _rich(meta_xlsx)
    infos = []
    for _ in stream:
        infos.append(stream.row_info)
    assert [i["row"] for i in infos] == list(range(len(infos)))
    assert infos[1]["height"] == 30 and infos[1]["custom_height"] is True
    assert infos[2]["hidden"] is True
    assert infos[3]["outline_level"] == 2
    assert infos[0]["hidden"] is False and infos[0]["height"] is None
    assert len(infos) == 20  # hidden empty row 25 is after the data


def test_row_info_for_cellless_rows(meta_xlsx):
    stream = _rich(meta_xlsx, sheet="Gaps")
    hidden = []
    for row in stream:
        hidden.append(stream.row_info["hidden"])
    assert hidden == [False, True, False]


def test_plain_stream_has_no_row_info(meta_xlsx):
    stream = CalamineWorkbook.from_path(meta_xlsx).stream_sheet_by_name("Meta")
    next(stream)
    assert stream.row_info is None


# ---- sheet level ----------------------------------------------------------------


def test_columns(meta_xlsx):
    cols = {(c["min"], c["max"]): c for c in _rich(meta_xlsx).column_dimensions}
    assert cols[(1, 1)]["width"] == 25 and cols[(1, 1)]["custom_width"] is True
    assert cols[(2, 2)]["hidden"] is True
    assert cols[(4, 5)]["outline_level"] == 1


def test_conditional_formats(meta_xlsx):
    cfs = _rich(meta_xlsx).conditional_formats
    by_type = {r["type"]: (cf["ranges"], r) for cf in cfs for r in cf["rules"]}
    ranges, cell_is = by_type["cellIs"]
    assert ranges == ["A1:A20"]
    assert cell_is["operator"] == "greaterThan"
    assert cell_is["formulas"] == ["10"]
    assert cell_is["style"]["font"]["bold"] is True
    assert cell_is["style"]["font"]["color"] == {"rgb": "FF9C0006"}
    assert cell_is["style"]["fill"]["bg_color"] == {"rgb": "FFFFC7CE"}

    _, scale = by_type["colorScale"]
    assert [v["type"] for v in scale["values"]] == ["min", "max"]
    assert [c["rgb"] for c in scale["colors"]] == ["FFF8696B", "FF63BE7B"]
    _, bar = by_type["dataBar"]
    assert bar["colors"][0]["rgb"] == "FF638EC6"
    _, icons = by_type["iconSet"]
    assert icons["icon_set"] == "3TrafficLights1"
    assert [v["value"] for v in icons["values"]] == ["0", "33", "67"]


def test_data_validations(meta_xlsx):
    dvs = {tuple(v["ranges"]): v for v in _rich(meta_xlsx).data_validations}
    lst = dvs[("F1:F10",)]
    assert lst["type"] == "list" and lst["formula1"] == '"Yes,No"'
    assert lst["allow_blank"] is True and lst["in_cell_dropdown"] is True
    assert lst["prompt"] == "Pick one" and lst["prompt_title"] == "Answer"
    whole = dvs[("G1",)]
    assert whole["type"] == "whole" and whole["operator"] == "between"
    assert (whole["formula1"], whole["formula2"]) == ("1", "10")
    assert whole["error"] == "1 to 10"


def test_sheet_settings(meta_xlsx):
    s = _rich(meta_xlsx).sheet_settings
    assert s["freeze_panes"] == "B2"
    assert (s["frozen_rows"], s["frozen_columns"]) == (1, 1)
    assert s["auto_filter"] == "A1:B20"
    assert s["protection"]["sheet"] is True
    assert s["protection"]["sort"] is False
    assert s["page_setup"]["orientation"] == "landscape"
    assert s["page_setup"]["paper_size"] == 9
    assert s["page_margins"]["left"] == 0.5
    assert s["print_options"]["grid_lines"] is True
    assert s["tab_color"] == {"rgb": "FF1072BA"}
    assert s["sheet_view"]["show_grid_lines"] is False


def test_sheet_settings_defaults():
    from pathlib import Path

    stream = CalamineWorkbook.from_path(
        Path(__file__).parent / "data" / "base.xlsx"
    ).stream_sheet_by_index(0, rich=True)
    s = stream.sheet_settings
    assert (
        s["freeze_panes"] is None
        and s["protection"] is None
        and s["auto_filter"] is None
    )
    assert (
        stream.hyperlinks == []
        and stream.conditional_formats == []
        and stream.data_validations == []
    )


# ---- ranges ---------------------------------------------------------------------


@pytest.mark.parametrize("rich", [False, True])
def test_row_and_column_range(meta_xlsx, rich):
    wb = CalamineWorkbook.from_path(meta_xlsx)
    full = list(wb.stream_sheet_by_name("Meta"))
    part = list(
        wb.stream_sheet_by_name(
            "Meta", rich=rich, min_row=5, max_row=9, min_col=1, max_col=2
        )
    )
    values = [[c.value for c in r] if rich else r for r in part]
    assert values == [r[1:3] for r in full[5:10]]
    if rich:
        assert (part[0][0].row, part[0][0].column) == (5, 1)


def test_range_single_cell_and_padding(meta_xlsx):
    wb = CalamineWorkbook.from_path(meta_xlsx)
    assert list(
        wb.stream_sheet_by_name("Meta", min_row=2, max_row=2, min_col=2, max_col=2)
    ) == [[42]]
    # max_col beyond the data pads with "", like openpyxl's iter_rows(max_col=...).
    assert list(wb.stream_sheet_by_name("Meta", max_row=0, min_col=6, max_col=8)) == [
        ["", "", ""]
    ]
    # Ranges past the data yield nothing.
    assert list(wb.stream_sheet_by_name("Meta", min_row=500)) == []


def test_range_stops_early_and_frees_stream(meta_xlsx):
    wb = CalamineWorkbook.from_path(meta_xlsx)
    s = wb.stream_sheet_by_name("Meta", max_row=1)
    assert len(list(s)) == 2
    assert (
        list(wb.stream_sheet_by_name("Meta", min_row=19)) != []
    )  # workbook still usable


def test_range_validation(meta_xlsx):
    wb = CalamineWorkbook.from_path(meta_xlsx)
    with pytest.raises(ValueError):
        wb.stream_sheet_by_name("Meta", min_row=5, max_row=4)
    with pytest.raises(OverflowError):
        wb.stream_sheet_by_name("Meta", min_row=-1)


def test_rich_text_shared_strings(tmp_path):
    # XlsxWriter stores rich strings in the shared strings table, like Excel.
    xlsxwriter = pytest.importorskip("xlsxwriter")
    path = tmp_path / "sst.xlsx"
    wb = xlsxwriter.Workbook(str(path))
    ws = wb.add_worksheet("S")
    bold = wb.add_format({"bold": True, "font_color": "red"})
    italic = wb.add_format({"italic": True, "font_size": 14})
    ws.write_rich_string("A1", "a ", bold, "bold", " and ", italic, "italic")
    ws.write("A2", "plain")
    ws.write_rich_string(
        "A3", "a ", bold, "bold", " and ", italic, "italic"
    )  # same string again
    wb.close()

    rows = list(CalamineWorkbook.from_path(path).stream_sheet_by_name("S", rich=True))
    a1, a2, a3 = rows[0][0], rows[1][0], rows[2][0]
    assert a1.value == "a bold and italic"
    assert [r["text"] for r in a1.rich_text] == ["a ", "bold", " and ", "italic"]
    assert a1.rich_text[1]["font"]["bold"] is True
    assert a1.rich_text[1]["font"]["color"] == {"rgb": "FFFF0000"}
    assert a1.rich_text[3]["font"]["italic"] is True
    assert a1.rich_text[3]["font"]["size"] == 14
    assert a2.rich_text is None
    assert a3.rich_text is a1.rich_text  # converted once per shared string
