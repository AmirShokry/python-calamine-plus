from datetime import date

import pytest
from python_calamine_plus import CalamineCell, CalamineWorkbook

openpyxl = pytest.importorskip("openpyxl")


@pytest.fixture(scope="module")
def rich_xlsx(tmp_path_factory):
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rich"

    ws["A1"] = "Header"
    ws["A1"].font = Font(
        name="Arial", size=14, bold=True, italic=True, color="FFFF0000"
    )
    ws["A1"].fill = PatternFill("solid", fgColor="FFFFFF00")
    ws["A1"].border = Border(bottom=Side(style="thin", color="FF0000FF"))
    ws["A1"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    ws.merge_cells("A1:C1")

    ws["A2"] = 1
    ws["B2"] = 0.25
    ws["B2"].number_format = "0.00%"
    ws["C2"] = date(2024, 1, 31)
    ws["C2"].number_format = "yyyy-mm-dd"
    ws["D2"].fill = PatternFill("solid", fgColor="FF00FF00")  # styled, no value

    ws["A3"] = "x"
    ws["B3"] = "commented"
    ws["B3"].comment = Comment("Check this", "Amir")
    ws["E3"].comment = Comment("On an empty cell", "Bo")  # empty cell in a value row
    ws.merge_cells("A4:A6")
    ws["A4"] = "tall"
    ws["A7"] = "after merge"  # so the merged rows 5-6 are inside the data
    ws["B7"] = "=SUM(A2:B2)*2"
    ws["C7"] = '=CONCATENATE("a","b")'
    ws["A9"] = "=A7"  # formula-only row: openpyxl saves no cached results

    ws["A20"].comment = Comment("Below the data", "Cy")  # row without values
    for col in "ABC":
        ws[f"{col}30"].fill = PatternFill(
            "solid", fgColor="FF00FFFF"
        )  # trailing styled rows

    path = tmp_path_factory.mktemp("rich") / "rich.xlsx"
    wb.save(path)
    return path


def _stream(path, rich):
    return CalamineWorkbook.from_path(path).stream_sheet_by_name("Rich", rich=rich)


def test_rich_rows_match_plain_rows(rich_xlsx):
    plain = list(_stream(rich_xlsx, rich=False))
    rich = list(_stream(rich_xlsx, rich=True))
    # Trailing styled-only rows are never streamed; the formula-only row 9 (no
    # saved result) is streamed in rich mode only.
    assert len(plain) == 7
    assert len(rich) == 9
    assert [c.value for c in rich[7]] == [""] * len(rich[7])
    assert rich[8][0].formula == "=A7" and rich[8][0].value == ""
    for prow, rrow in zip(plain, rich):
        assert all(isinstance(c, CalamineCell) for c in rrow)
        values = [c.value for c in rrow]
        # Rich rows may be wider (styled empty cells such as D2), never different.
        assert values[: len(prow)] == prow
        assert all(v == "" for v in values[len(prow) :])


def test_rich_positions(rich_xlsx):
    for r, row in enumerate(_stream(rich_xlsx, rich=True)):
        for c, cell in enumerate(row):
            assert (cell.row, cell.column) == (r, c)


def test_rich_styles(rich_xlsx):
    rows = list(_stream(rich_xlsx, rich=True))
    a1 = rows[0][0].style
    assert a1["font"]["name"] == "Arial"
    assert a1["font"]["size"] == 14
    assert a1["font"]["bold"] and a1["font"]["italic"]
    assert a1["font"]["color"] == {"rgb": "FFFF0000"}
    assert a1["fill"]["pattern"] == "solid"
    assert a1["fill"]["fg_color"]["rgb"] == "FFFFFF00"
    assert a1["border"]["bottom"] == {"style": "thin", "color": {"rgb": "FF0000FF"}}
    assert a1["border"]["top"]["style"] is None
    assert a1["alignment"]["horizontal"] == "center"
    assert a1["alignment"]["vertical"] == "top"
    assert a1["alignment"]["wrap_text"] is True

    b2, c2, d2 = rows[1][1], rows[1][2], rows[1][3]
    assert b2.value == 0.25 and b2.style["number_format"] == "0.00%"
    assert c2.value == date(2024, 1, 31) and c2.style["number_format"] == "yyyy-mm-dd"
    assert d2.value == "" and d2.style["fill"]["fg_color"]["rgb"] == "FF00FF00"

    plain = rows[2][0]
    assert plain.style_id == 0
    assert plain.style["font"]["bold"] is False


def test_rich_styles_shared(rich_xlsx):
    stream = _stream(rich_xlsx, rich=True)
    rows = list(stream)
    styles = stream.styles
    a1 = rows[0][0]
    assert styles[a1.style_id] is a1.style  # one dict per format, shared by cells
    assert rows[2][0].style is rows[3][0].style


def test_rich_comments(rich_xlsx):
    stream = _stream(rich_xlsx, rich=True)
    rows = list(stream)
    assert rows[2][1].comment == {"author": "Amir", "text": "Check this"}
    assert rows[2][4].value == ""
    assert rows[2][4].comment == {"author": "Bo", "text": "On an empty cell"}
    assert rows[0][0].comment is None
    # Comments outside the streamed rows are still available on the stream.
    assert stream.comments[(19, 0)] == {"author": "Cy", "text": "Below the data"}
    assert len(stream.comments) == 3


def test_rich_merged(rich_xlsx):
    stream = _stream(rich_xlsx, rich=True)
    rows = list(stream)
    assert sorted(stream.merged_cell_ranges) == [((0, 0), (0, 2)), ((3, 0), (5, 0))]
    assert [c.merged_range for c in rows[0][:4]] == [((0, 0), (0, 2))] * 3 + [None]
    for r in (3, 4, 5):
        assert rows[r][0].merged_range == ((3, 0), (5, 0))
        assert rows[r][1].merged_range is None
    assert rows[1][0].merged_range is None
    assert rows[6][0].merged_range is None and rows[6][0].value == "after merge"


def test_rich_matches_openpyxl_styles(rich_xlsx):
    ox = openpyxl.load_workbook(rich_xlsx)["Rich"]
    for row in _stream(rich_xlsx, rich=True):
        for cell in row:
            ref = ox.cell(row=cell.row + 1, column=cell.column + 1)
            assert cell.style["number_format"] == ref.number_format
            assert cell.style["font"]["bold"] == bool(ref.font.b)
            assert cell.style["fill"]["pattern"] == ref.fill.patternType


def test_plain_stream_has_no_rich_info(rich_xlsx):
    stream = _stream(rich_xlsx, rich=False)
    assert stream.rich is False
    assert stream.styles is None
    assert stream.comments is None
    assert stream.merged_cell_ranges is None
    assert isinstance(next(stream)[0], str)


def test_rich_on_sheet_without_extras():
    from pathlib import Path

    path = Path(__file__).parent / "data" / "base.xlsx"
    wb = CalamineWorkbook.from_path(path)
    stream = wb.stream_sheet_by_index(0, rich=True)
    rows = list(stream)
    assert [[c.value for c in r] for r in rows] == list(wb.stream_sheet_by_index(0))
    assert stream.comments == {}
    assert stream.merged_cell_ranges == []


def test_rich_formulas(rich_xlsx):
    rows = list(_stream(rich_xlsx, rich=True))
    assert rows[6][1].formula == "=SUM(A2:B2)*2"
    assert rows[6][2].formula == '=CONCATENATE("a","b")'
    assert rows[6][1].value == ""  # openpyxl never calculates, so no saved result
    assert rows[6][0].formula is None and rows[6][0].value == "after merge"
    assert all(c.formula is None for r in rows[:6] for c in r)


def test_rich_formulas_with_cached_values():
    from pathlib import Path

    path = Path(__file__).parent / "data" / "base.xlsx"
    wb = CalamineWorkbook.from_path(path)
    for name in wb.sheet_names:
        for row in wb.stream_sheet_by_name(name, rich=True):
            for cell in row:
                assert cell.formula is None or cell.formula.startswith("=")
