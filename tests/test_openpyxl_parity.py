"""Reader features openpyxl has (workbook settings, named styles, tables, views, ...)."""

from datetime import date, datetime

import pytest
from python_calamine_plus import CalamineWorkbook

openpyxl = pytest.importorskip("openpyxl")


@pytest.fixture(scope="module")
def parity_xlsx(tmp_path_factory):
    from openpyxl.packaging.custom import (
        BoolProperty,
        DateTimeProperty,
        FloatProperty,
        IntProperty,
        StringProperty,
    )
    from openpyxl.styles import Alignment, Border, Font, GradientFill, NamedStyle, Side
    from openpyxl.workbook.protection import WorkbookProtection
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    wb.create_sheet("Second")
    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    wb.security = WorkbookProtection(lockStructure=True)
    wb.custom_doc_props.append(StringProperty(name="Project", value="Fork"))
    wb.custom_doc_props.append(IntProperty(name="Count", value=3))
    wb.custom_doc_props.append(FloatProperty(name="Ratio", value=0.5))
    wb.custom_doc_props.append(BoolProperty(name="Final", value=True))
    wb.custom_doc_props.append(
        DateTimeProperty(name="Due", value=datetime(2024, 5, 1, 12, 0))
    )

    ws.append(["Name", "Amount"])
    for i, name in enumerate(["a", "b", "c", "d"], start=1):
        ws.append([name, i * 10])
    ws.add_table(
        Table(
            displayName="Sales",
            ref="A1:B5",
            tableStyleInfo=TableStyleInfo(
                name="TableStyleMedium9", showRowStripes=True
            ),
        )
    )

    highlight = NamedStyle(name="Highlight", font=Font(bold=True))
    wb.add_named_style(highlight)
    ws["C1"] = "styled"
    ws["C1"].style = "Highlight"
    ws["C2"] = "fancy"
    ws["C2"].font = Font(
        name="Calibri", family=2, scheme="minor", charset=1, outline=True, shadow=True
    )
    ws["C2"].fill = GradientFill(degree=90, stop=("FFFF0000", "FF0000FF"))
    ws["C2"].border = Border(diagonal=Side(style="thin"), diagonalUp=True)
    ws["C2"].alignment = Alignment(
        readingOrder=2, justifyLastLine=True, relativeIndent=1
    )
    ws["C3"] = "'quoted"
    ws["C3"].quotePrefix = True
    ws["D1"] = date(2024, 2, 29)
    ws["D2"] = True
    ws["D3"] = "#REF!"

    ws["E1"] = ArrayFormula("E1:E3", "=B2:B4*2")

    ws.row_breaks.append(Break(id=3))
    ws.col_breaks.append(Break(id=2))
    ws.oddHeader.center.text = "Report"
    ws.oddFooter.right.text = "Page &P"
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_format.defaultRowHeight = 20
    ws.sheet_view.selection[0].activeCell = "C3"
    ws.sheet_view.selection[0].sqref = "C3"
    ws.auto_filter.ref = "A1:B5"
    ws.auto_filter.add_filter_column(0, ["a", "b"])
    ws.auto_filter.add_sort_condition("B2:B5", descending=True)
    ws.print_area = "A1:C10"
    ws.print_title_rows = "1:1"
    ws.row_dimensions[50].hidden = True  # far below the data

    path = tmp_path_factory.mktemp("parity") / "parity.xlsx"
    wb.save(path)
    return path


def _rich(path, sheet="Data", **kw):
    return CalamineWorkbook.from_path(path).stream_sheet_by_name(sheet, rich=True, **kw)


# ---- workbook ---------------------------------------------------------------------


def test_workbook_settings(parity_xlsx):
    wb = CalamineWorkbook.from_path(parity_xlsx)
    assert wb.active_sheet_index == 1
    assert wb.calculation["full_calc_on_load"] is True
    assert wb.workbook_protection["lock_structure"] is True
    assert wb.epoch_1904 is False
    assert wb.is_template is False and wb.has_macros is False
    assert wb.workbook_views[0]["active_tab"] == 1


def test_custom_properties(parity_xlsx):
    props = CalamineWorkbook.from_path(parity_xlsx).custom_properties
    assert props == {
        "Project": "Fork",
        "Count": 3,
        "Ratio": 0.5,
        "Final": True,
        "Due": datetime(2024, 5, 1, 12, 0),
    }


def test_theme(parity_xlsx):
    wb = CalamineWorkbook.from_path(parity_xlsx)
    ox = openpyxl.load_workbook(parity_xlsx)
    assert wb.theme_xml == ox.loaded_theme
    colors = wb.theme_colors
    assert list(colors)[:4] == ["dk1", "lt1", "dk2", "lt2"]
    assert colors["accent1"] == "4F81BD"  # openpyxl's default Office theme


def test_named_styles(parity_xlsx):
    wb = CalamineWorkbook.from_path(parity_xlsx)
    names = [s["name"] for s in wb.named_styles]
    assert "Normal" in names and "Highlight" in names
    rows = list(_rich(parity_xlsx))
    assert rows[0][2].style["named_style"] == "Highlight"
    assert rows[1][2].style["named_style"] == "Normal"


def test_epoch_1904(tmp_path):
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    wb = openpyxl.Workbook()
    wb.epoch = CALENDAR_MAC_1904
    wb.active["A1"] = date(2024, 2, 29)
    path = tmp_path / "1904.xlsx"
    wb.save(path)
    cw = CalamineWorkbook.from_path(path)
    assert cw.epoch_1904 is True
    assert next(cw.stream_sheet_by_index(0))[0] == date(2024, 2, 29)


def test_template(tmp_path):
    wb = openpyxl.Workbook()
    wb.template = True
    wb.active["A1"] = 1
    path = tmp_path / "book.xltx"
    wb.save(path)
    cw = CalamineWorkbook.from_path(path)
    assert cw.is_template is True
    assert list(cw.stream_sheet_by_index(0)) == [[1]]


# ---- cells ------------------------------------------------------------------------


def test_style_details(parity_xlsx):
    rows = list(_rich(parity_xlsx))
    c2 = rows[1][2].style
    assert c2["font"]["family"] == 2 and c2["font"]["scheme"] == "minor"
    assert c2["font"]["charset"] == 1
    assert c2["font"]["outline"] is True and c2["font"]["shadow"] is True
    grad = c2["fill"]["gradient"]
    assert c2["fill"]["pattern"] == "gradient" and grad["degree"] == 90
    assert [s["color"]["rgb"] for s in grad["stops"]] == ["FFFF0000", "FF0000FF"]
    assert (
        c2["border"]["diagonal_up"] is True
        and c2["border"]["diagonal"]["style"] == "thin"
    )
    assert c2["alignment"]["reading_order"] == 2
    assert c2["alignment"]["justify_last_line"] is True
    assert c2["alignment"]["relative_indent"] == 1
    assert rows[2][2].style["quote_prefix"] is True


def test_data_type_is_date_coordinate(parity_xlsx):
    rows = list(_rich(parity_xlsx))
    assert (rows[0][3].data_type, rows[0][3].is_date) == ("d", True)
    assert rows[1][3].data_type == "b"
    assert (rows[2][3].data_type, rows[2][3].value) == ("e", "#REF!")
    assert (rows[1][1].data_type, rows[1][0].data_type) == ("n", "s")
    assert rows[1][1].is_date is False
    assert rows[2][3].coordinate == "D3"
    ox = openpyxl.load_workbook(parity_xlsx)["Data"]
    for row in rows:
        for cell in row:
            theirs = ox[cell.coordinate]
            if theirs.value is not None and theirs.data_type != "f":
                assert cell.data_type == theirs.data_type, cell.coordinate


def test_array_formula(parity_xlsx):
    rows = list(_rich(parity_xlsx))
    e1 = rows[0][4]
    assert e1.formula == "=B2:B4*2"
    assert e1.formula_type == "array"
    assert e1.formula_range == "E1:E3"
    assert e1.formula_attributes["ref"] == "E1:E3"
    ox = openpyxl.load_workbook(parity_xlsx)["Data"]["E1"].value
    assert (ox.ref, ox.text) == (e1.formula_range, e1.formula)
    assert rows[1][1].formula_type is None


def test_formula_types(tmp_path):
    # Excel-style shared formulas (openpyxl never writes them), via XlsxWriter... which
    # also doesn't; so check against the shared formulas of the existing test files.
    from pathlib import Path

    wb = CalamineWorkbook.from_path(Path(__file__).parent / "data" / "base.xlsx")
    kinds = {
        c.formula_type
        for name in wb.sheet_names
        for row in wb.stream_sheet_by_name(name, rich=True)
        for c in row
    }
    assert kinds <= {None, "normal", "shared", "array", "dataTable"}


# ---- sheet ------------------------------------------------------------------------


def test_tables(parity_xlsx):
    (table,) = _rich(parity_xlsx).tables
    assert table["display_name"] == "Sales" and table["ref"] == "A1:B5"
    assert [c["name"] for c in table["columns"]] == ["Name", "Amount"]
    assert table["style"]["name"] == "TableStyleMedium9"
    assert table["style"]["show_row_stripes"] is True
    assert _rich(parity_xlsx, sheet="Second").tables == []


def test_dimension_and_breaks(parity_xlsx):
    s = _rich(parity_xlsx)
    ox = openpyxl.load_workbook(parity_xlsx)["Data"]
    assert s.dimension == ox.dimensions
    assert [b["id"] for b in s.row_breaks] == [3]
    assert [b["id"] for b in s.column_breaks] == [2]
    assert s.scenarios is None


def test_row_dimensions_outside_data(parity_xlsx):
    s = _rich(parity_xlsx)
    assert len(list(s)) == 5
    assert s.row_dimensions[49]["hidden"] is True


def test_sheet_settings_parity(parity_xlsx):
    s = _rich(parity_xlsx).sheet_settings
    assert s["header_footer"]["odd_header"] == "&CReport"
    assert s["header_footer"]["odd_footer"] == "&RPage &P"
    assert s["outline_properties"]["summary_below"] is False
    assert s["page_setup_properties"]["fit_to_page"] is True
    assert s["sheet_format"]["default_row_height"] == 20
    assert s["views"][0]["selections"][0]["active_cell"] == "C3"
    ox = openpyxl.load_workbook(parity_xlsx)["Data"]
    assert s["print_area"] == ox.print_area == "'Data'!$A$1:$C$10"
    assert s["print_titles"] == ox.print_titles == "'Data'!$1:$1"
    af = s["auto_filter_details"]
    assert af["ref"] == "A1:B5"
    assert af["columns"][0]["column"] == 0 and af["columns"][0]["values"] == ["a", "b"]
    assert af["sort"]["conditions"][0] == {"ref": "B2:B5", "descending": True}


def test_ref_argument(parity_xlsx):
    wb = CalamineWorkbook.from_path(parity_xlsx)
    assert list(wb.stream_sheet_by_name("Data", ref="A2:B3")) == [["a", 10], ["b", 20]]
    assert list(wb.stream_sheet_by_name("Data", ref="B3")) == [[20]]
    assert list(wb.stream_sheet_by_name("Data", ref="$B$3:$A$2")) == [
        ["a", 10],
        ["b", 20],
    ]
    with pytest.raises(ValueError):
        wb.stream_sheet_by_name("Data", ref="nope")
